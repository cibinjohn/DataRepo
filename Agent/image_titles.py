from pptx.enum.shapes import MSO_SHAPE_TYPE

def get_shape_bounds(shape):
    return {
        "left": shape.left,
        "top": shape.top,
        "right": shape.left + shape.width,
        "bottom": shape.top + shape.height,
        "center_x": shape.left + shape.width / 2,
    }

def link_labels_to_images(slide):
    labels = []
    images = []

    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images.append(shape)
        elif getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:  # skip empty text boxes
                labels.append(shape)

    associations = []

    for label in labels:
        lb = get_shape_bounds(label)
        best_match = None
        best_distance = float("inf")

        for img in images:
            ib = get_shape_bounds(img)

            # Does the label's horizontal center fall within the image's horizontal span?
            # (with some tolerance, since labels are often narrower than the image group)
            horizontal_overlap = ib["left"] - 100000 <= lb["center_x"] <= ib["right"] + 100000

            if not horizontal_overlap:
                continue

            # Vertical distance from label to image (label should be above or beside it)
            vertical_gap = abs(ib["top"] - lb["top"])

            if vertical_gap < best_distance:
                best_distance = vertical_gap
                best_match = img

        associations.append({
            "label_text": label.text_frame.text.strip(),
            "image_shape_id": best_match.shape_id if best_match else None,
            "image_name": best_match.name if best_match else None,
        })

    return associations

# Usage
second_slide = presentation.slides[1]
result = link_labels_to_images(second_slide)
for r in result:
    print(r)

########
import matplotlib.pyplot as plt
from PIL import Image
import io

def display_label_image_pairs(associations, slide):
    # Build a lookup from shape_id -> shape for quick access
    shape_lookup = {shape.shape_id: shape for shape in slide.shapes}

    n = len(associations)
    fig, axes = plt.subplots(1, n, figsize=(6 * n, 6))

    if n == 1:
        axes = [axes]  # keep iterable if only one pair

    for ax, assoc in zip(axes, associations):
        label = assoc["label_text"]
        image_id = assoc["image_shape_id"]

        if image_id is None:
            ax.text(0.5, 0.5, f"No image matched\nfor '{label}'",
                     ha="center", va="center")
            ax.axis("off")
            continue

        shape = shape_lookup[image_id]
        image_bytes = shape.image.blob
        img = Image.open(io.BytesIO(image_bytes))

        ax.imshow(img)
        ax.set_title(label, fontsize=14, fontweight="bold")
        ax.axis("off")

    plt.tight_layout()
    plt.show()

# Usage
second_slide = presentation.slides[1]
result = link_labels_to_images(second_slide)
display_label_image_pairs(result, second_slide)
